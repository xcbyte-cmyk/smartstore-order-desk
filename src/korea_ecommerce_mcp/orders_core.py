"""SmartStore order collection and a durable, preview-first shipping ledger."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import secrets
import sqlite3
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlencode

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from korea_ecommerce_mcp.adapters.naver import NaverAdapter
from korea_ecommerce_mcp.config import Settings

KST = timezone(timedelta(hours=9))
CARRIERS = {
    "CJGLS": "CJ대한통운",
    "HANJIN": "한진택배",
    "KGB": "로젠택배",
    "EPOST": "우체국택배",
    "HYUNDAI": "롯데택배",
    "KDEXP": "경동택배",
}
PREFIX = "/v1/pay-order/seller/product-orders"


def now():
    return datetime.now(KST)


def stamp():
    return now().isoformat(timespec="seconds")


def account_key(settings):
    return hashlib.sha256(
        f"{settings.naver_client_id}:{settings.naver_account_id or 'SELF'}".encode()
    ).hexdigest()


def public_error(exc):
    if isinstance(exc, httpx.HTTPStatusError):
        return (
            f"네이버 API 오류 (HTTP {exc.response.status_code}). API 권한과 허용 IP를 확인하세요."
        )
    if isinstance(exc, httpx.RequestError):
        return "네이버 연결에 실패했습니다. 네트워크를 확인해 주세요."
    if isinstance(exc, ValueError):
        return str(exc)
    return "처리 중 오류가 발생했습니다. 설정 또는 파일 내용을 확인해 주세요."


class OrderDB:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS orders (
                id TEXT PRIMARY KEY, data TEXT NOT NULL, updated TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY, at TEXT, kind TEXT, message TEXT);
            CREATE TABLE IF NOT EXISTS uploads (token TEXT PRIMARY KEY, created TEXT, data TEXT);
            CREATE TABLE IF NOT EXISTS shipments (
                id TEXT PRIMARY KEY, tracking TEXT, carrier TEXT,
                state TEXT, message TEXT, updated TEXT);
            """)
            # A process exit after sending must never turn into an automatic retry.
            db.execute(
                "UPDATE shipments SET state='unknown',"
                "message='송신 중 종료. 네이버 상태를 재조회하세요.' WHERE state='sending'"
            )

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def get(self, key, default=None):
        with self.connect() as db:
            row = db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
            return json.loads(row[0]) if row else default

    def set(self, key, value):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO settings VALUES (?,?)", (key, json.dumps(value)))

    def event(self, kind, message):
        with self.connect() as db:
            db.execute(
                "INSERT INTO events(at,kind,message) VALUES (?,?,?)", (stamp(), kind, message)
            )
            db.execute(
                "DELETE FROM events WHERE id NOT IN "
                "(SELECT id FROM events ORDER BY id DESC LIMIT 500)"
            )

    def events(self):
        with self.connect() as db:
            return [dict(r) for r in db.execute("SELECT * FROM events ORDER BY id DESC LIMIT 30")]

    def upsert(self, orders):
        with self.connect() as db:
            db.executemany(
                "INSERT OR REPLACE INTO orders VALUES (?,?,?)",
                [(o["id"], json.dumps(o, ensure_ascii=False), stamp()) for o in orders],
            )

    def orders(self):
        with self.connect() as db:
            rows = [
                json.loads(r[0])
                for r in db.execute("SELECT data FROM orders ORDER BY updated DESC")
            ]
            shipments = {r["id"]: dict(r) for r in db.execute("SELECT * FROM shipments")}
        for row in rows:
            row["shipment"] = shipments.get(row["id"])
        return rows

    def shipment(self, oid):
        with self.connect() as db:
            row = db.execute("SELECT * FROM shipments WHERE id=?", (oid,)).fetchone()
            return dict(row) if row else None

    def shipping(self, row, state, message):
        with self.connect() as db:
            db.execute(
                "INSERT OR REPLACE INTO shipments VALUES (?,?,?,?,?,?)",
                (row["id"], row["tracking"], row["carrier"], state, message, stamp()),
            )

    def preview(self, rows, account):
        token = secrets.token_urlsafe(32)
        payload = {"rows": rows, "account": account}
        with self.connect() as db:
            db.execute(
                "DELETE FROM uploads WHERE created < ?",
                ((now() - timedelta(minutes=15)).isoformat(),),
            )
            db.execute("INSERT INTO uploads VALUES (?,?,?)", (token, stamp(), json.dumps(payload)))
        return token

    def consume(self, token, account):
        with self.connect() as db:
            row = db.execute("SELECT * FROM uploads WHERE token=?", (token,)).fetchone()
            if not row or now() - datetime.fromisoformat(row["created"]) > timedelta(minutes=15):
                raise ValueError(
                    "업로드 검토가 만료되었거나 이미 송신했습니다. 파일을 다시 올려 주세요."
                )
            data = json.loads(row["data"])
            if data["account"] != account:
                raise ValueError("연결 계정이 바뀌었습니다. 파일을 다시 올려 주세요.")
            db.execute("DELETE FROM uploads WHERE token=?", (token,))
        return data["rows"]


def normalize(detail):
    p, o = detail.get("productOrder", {}), detail.get("order", {})
    a, delivery = p.get("shippingAddress") or {}, detail.get("delivery") or {}
    if not p.get("productOrderId"):
        raise ValueError("상품주문번호가 없는 API 응답입니다. 수집 완료로 기록하지 않았습니다.")
    return {
        "id": str(p["productOrderId"]),
        "order_id": str(o.get("orderId", "")),
        "date": o.get("paymentDate") or o.get("orderDate") or "",
        "name": p.get("productName", ""),
        "option": p.get("productOption", ""),
        "quantity": p.get("remainQuantity", p.get("quantity", 0)),
        "recipient": a.get("name", ""),
        "phone": a.get("tel1", ""),
        "postcode": a.get("zipCode", ""),
        "address": " ".join(
            str(a.get(k) or "") for k in ["baseAddress", "detailedAddress"]
        ).strip(),
        "memo": p.get("shippingMemo", ""),
        "status": p.get("productOrderStatus", ""),
        "place_status": p.get("placeOrderStatus", ""),
        "claim": p.get("claimStatus") or "",
        "method": p.get("expectedDeliveryMethod") or "",
        "tracking": delivery.get("trackingNumber") or "",
        "carrier": delivery.get("deliveryCompany") or delivery.get("deliveryCompanyCode") or "",
    }


class NaverOrders:
    def __init__(self, settings):
        self.adapter = NaverAdapter(settings)

    async def close(self):
        await self.adapter.close()

    async def request(self, method, suffix, body=None, params=None):
        path = PREFIX + suffix
        if params:
            path += "?" + urlencode(params)
        # Reads and writes are issued once. No retry of ambiguous dispatches.
        token = await self.adapter._token()
        response = await self.adapter.client.request(
            method,
            self.adapter.base_url + path,
            headers={"Authorization": f"Bearer {token}"},
            json=body,
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("code"):
            raise ValueError(
                "네이버에서 요청을 거절했습니다. API 설정과 주문 상태를 확인해 주세요."
            )
        return payload.get("data")

    async def details(self, ids):
        rows = []
        for start in range(0, len(ids), 300):
            batch = ids[start : start + 300]
            data = await self.request("POST", "/query", {"productOrderIds": batch})
            if not isinstance(data, list):
                raise ValueError("주문 상세 응답이 올바르지 않습니다.")
            normalized = [normalize(d) for d in data]
            if set(batch) != {r["id"] for r in normalized}:
                raise ValueError("일부 주문 상세가 누락되었습니다. 다시 수집해 주세요.")
            rows.extend(normalized)
            await asyncio.sleep(0.55)
        return rows

    async def collect_window(self, begin, end):
        params = {
            "lastChangedFrom": begin.isoformat(timespec="milliseconds"),
            "lastChangedTo": end.isoformat(timespec="milliseconds"),
            "limitCount": 300,
        }
        ids, seen = set(), set()
        while True:
            data = await self.request("GET", "/last-changed-statuses", params=params)
            if data is None:
                break
            if not isinstance(data, dict):
                raise ValueError("변경 주문 응답 형식이 올바르지 않습니다.")
            ids.update(str(r["productOrderId"]) for r in data.get("lastChangeStatuses", []))
            more = data.get("more")
            if not more:
                break
            cursor = (more["moreFrom"], more["moreSequence"])
            if cursor in seen:
                raise ValueError(
                    "네이버의 다음 페이지가 반복됩니다. 수집 완료로 기록하지 않았습니다."
                )
            seen.add(cursor)
            params.update(lastChangedFrom=cursor[0], moreSequence=cursor[1])
            await asyncio.sleep(0.55)
        return await self.details(sorted(ids))

    async def dispatch(self, rows):
        return await self.request(
            "POST",
            "/dispatch",
            {
                "dispatchProductOrders": [
                    {
                        "productOrderId": r["id"],
                        "deliveryMethod": "DELIVERY",
                        "deliveryCompanyCode": r["carrier"],
                        "trackingNumber": r["tracking"],
                        "dispatchDate": now().isoformat(timespec="milliseconds"),
                    }
                    for r in rows
                ]
            },
        )


def shipping_problem(order):
    if not order:
        return "수집된 주문이 없습니다. 주문을 먼저 수집해 주세요."
    if order["status"] != "PAYED":
        return "결제완료 주문만 송신할 수 있습니다. 현재 상태: " + order["status"]
    if order["claim"]:
        return "취소·반품·교환 내역 확인이 필요합니다."
    if order["method"] != "DELIVERY":
        return "일반 택배 배송 주문만 지원합니다."
    if order["place_status"] != "OK":
        return "스마트스토어에서 발주 확인 후 주문을 다시 수집해 주세요."
    return ""


def workbook_bytes(headers, rows, title="발주서", template=False):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cells in sheet.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                # Prevent spreadsheet formulas and preserve identifiers/leading zeroes.
                cell.data_type = "s"
                cell.number_format = "@"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173D35")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = (
            42 if header in ["주소", "상품명", "배송메모"] else 22
        )
    if template:
        valid = DataValidation(type="list", formula1='"' + ",".join(CARRIERS.values()) + '"')
        sheet.add_data_validation(valid)
        valid.add("B2:B1001")
        for col in ["A", "B", "C"]:
            for index in range(2, 1002):
                sheet[f"{col}{index}"].number_format = "@"
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def parse_invoices(content, orders, db):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("파일은 5MB 이하로 올려 주세요.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("압축 해제 크기가 너무 큰 파일입니다.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("정상적인 .xlsx 파일을 올려 주세요.") from exc
    try:
        sheet = book.active
        if sheet is None or (sheet.max_row or 0) > 10001 or (sheet.max_column or 0) > 100:
            raise ValueError("양식 크기가 너무 큽니다. 송장 열과 최대 1,000건만 남겨 주세요.")
        iterator = sheet.iter_rows()
        headers = [str(c.value or "").strip() for c in next(iterator, [])]
        expected = ["상품주문번호", "택배사", "송장번호"]
        if any(headers.count(k) != 1 for k in expected):
            raise ValueError(
                "첫 행에 상품주문번호, 택배사, 송장번호 열이 각각 하나씩 있어야 합니다."
            )
        indexes = [headers.index(k) for k in expected]
        rows = []
        by_id = {o["id"]: o for o in orders}
        aliases = {v: k for k, v in CARRIERS.items()} | {k: k for k in CARRIERS}
        for row_number, cells in enumerate(iterator, 2):
            selected = [cells[i] if i < len(cells) else None for i in indexes]
            if not any(c and c.value is not None for c in selected):
                continue
            if len(rows) >= 1000:
                raise ValueError("한 번에 1,000행까지 업로드할 수 있습니다.")
            values = [str(c.value).strip() if c and c.value is not None else "" for c in selected]
            oid, carrier, tracking = values
            error = ""
            if any(c and c.data_type == "f" for c in selected):
                error = "수식은 사용할 수 없습니다. 값을 텍스트로 입력해 주세요."
            elif any(c and isinstance(c.value, (float, int)) for c in [selected[0], selected[2]]):
                error = "상품주문번호·송장번호는 자릿수 손실을 막도록 텍스트 형식으로 입력하세요."
            elif not re.fullmatch(r"[0-9]{8,24}", oid):
                error = "상품주문번호 형식을 확인해 주세요."
            elif carrier not in aliases:
                error = "지원 택배사를 선택해 주세요: " + ", ".join(CARRIERS.values())
            elif not re.fullmatch(r"[0-9]{6,30}", tracking):
                error = "송장번호는 공백·하이픈 없는 숫자 6~30자리로 입력해 주세요."
            else:
                error = shipping_problem(by_id.get(oid))
            prior = db.shipment(oid)
            if (
                not error
                and prior
                and prior["state"] in ["success", "accepted", "unknown", "sending"]
            ):
                error = (
                    "이미 송신했거나 결과 확인 중인 주문입니다. 처리 내역에서 상태를 확인해 주세요."
                )
            rows.append(
                {
                    "row": row_number,
                    "id": oid,
                    "carrier": aliases.get(carrier, carrier),
                    "tracking": tracking,
                    "error": error,
                    "name": by_id.get(oid, {}).get("name", ""),
                }
            )
        counts = Counter(r["id"] for r in rows)
        for row in rows:
            if counts[row["id"]] > 1:
                row["error"] = "파일 안에 상품주문번호가 중복되어 있습니다."
        if not rows:
            raise ValueError("송장 데이터가 없습니다. 2행부터 입력해 주세요.")
        return rows
    finally:
        book.close()


class OrderService:
    def __init__(self, db, settings_factory=Settings, api_factory=NaverOrders):
        self.db, self.settings_factory, self.api_factory = db, settings_factory, api_factory
        self.lock = asyncio.Lock()
        self.busy = ""

    def settings(self):
        settings = self.settings_factory()
        if not settings.naver_client_id or not settings.naver_client_secret:
            raise ValueError("스마트스토어 API 연결 정보를 먼저 저장해 주세요.")
        bound = self.db.get("account")
        key = account_key(settings)
        if bound and bound != key:
            raise ValueError("수집 데이터와 연결 계정이 다릅니다. 기존 계정으로 연결해 주세요.")
        return settings

    async def collect(self, lookback_days=None):
        if self.lock.locked():
            raise ValueError("다른 작업이 진행 중입니다. 완료 후 다시 시도해 주세요.")
        async with self.lock:
            self.busy = "주문 수집 중"
            api = None
            count = 0
            try:
                settings = self.settings()
                api = self.api_factory(settings)
                end = now() - timedelta(seconds=5)
                cursor = self.db.get("cursor")
                begin = (
                    (end - timedelta(days=lookback_days))
                    if lookback_days
                    else (
                        datetime.fromisoformat(cursor) - timedelta(minutes=2)
                        if cursor
                        else end - timedelta(days=1)
                    )
                )
                if end - begin > timedelta(days=90):
                    raise ValueError(
                        "마지막 수집 후 90일이 지났습니다. 기간 재수집을 실행해 주세요."
                    )
                self.db.set("account", account_key(settings))
                while begin < end:
                    until = min(begin + timedelta(hours=23), end)
                    rows = await api.collect_window(begin, until)
                    self.db.upsert(rows)
                    count += len(rows)
                    # Commit the cursor only after all pages and details in the window are durable.
                    if not lookback_days:
                        self.db.set("cursor", until.isoformat())
                    begin = until
                # Some status changes do not appear in the change feed.
                pending = [
                    o["id"]
                    for o in self.db.orders()
                    if o["status"] == "PAYED"
                    or (o.get("shipment") and o["shipment"]["state"] in ["unknown", "accepted"])
                ]
                if pending:
                    self.db.upsert(await api.details(pending))
                self.reconcile()
                self.db.set("last_success", stamp())
                self.db.set("last_error", "")
                self.db.event("success", f"주문 수집 완료 · 변경 {count}건 갱신")
                return count
            except Exception as exc:
                message = public_error(exc)
                self.db.set("last_error", message)
                self.db.event("error", message)
                raise ValueError(message) from exc
            finally:
                self.db.set(
                    "next_run", (now() + timedelta(minutes=self.db.get("interval", 60))).isoformat()
                )
                self.busy = ""
                if api:
                    await api.close()

    def reconcile(self):
        for order in self.db.orders():
            prior = order.get("shipment")
            if not prior or prior["state"] not in ["accepted", "unknown"]:
                continue
            if (
                order["status"] in ["DELIVERING", "DELIVERED", "PURCHASE_DECIDED"]
                and order["tracking"] == prior["tracking"]
                and order["carrier"] == prior["carrier"]
            ):
                self.db.shipping(prior, "success", "네이버 주문 상세에서 송장 반영 확인")

    async def send(self, token):
        if self.lock.locked():
            raise ValueError("다른 작업이 진행 중입니다.")
        if not self.db.get("dispatch_enabled", False):
            raise ValueError("연결 설정에서 송장 송신을 활성화해 주세요.")
        async with self.lock:
            settings = self.settings()
            rows = self.db.consume(token, account_key(settings))
            valid = [r for r in rows if not r["error"]]
            if not valid:
                raise ValueError("송신 가능한 행이 없습니다.")
            self.busy = "송장 송신 중"
            api = self.api_factory(settings)
            try:
                fresh = await api.details([r["id"] for r in valid])
                self.db.upsert(fresh)
                by_id = {o["id"]: o for o in fresh}
                sendable = []
                for row in valid:
                    prior = self.db.shipment(row["id"])
                    if prior and prior["state"] in ["success", "accepted", "unknown", "sending"]:
                        continue
                    error = shipping_problem(by_id.get(row["id"]))
                    if error:
                        self.db.shipping(row, "failed", error)
                    else:
                        sendable.append(row)
                for offset in range(0, len(sendable), 30):
                    batch = sendable[offset : offset + 30]
                    for row in batch:
                        self.db.shipping(row, "sending", "송신 요청 중")
                    try:
                        result = await api.dispatch(batch)
                        if not isinstance(result, dict):
                            raise ValueError("발송 처리 응답이 없습니다.")
                        successes = set(result.get("successProductOrderIds") or [])
                        failures = {
                            str(r["productOrderId"]): r
                            for r in result.get("failProductOrderInfos") or []
                        }
                        for row in batch:
                            if row["id"] in successes:
                                self.db.shipping(
                                    row, "accepted", "네이버 처리 성공 응답 · 상세 반영 확인 대기"
                                )
                            elif row["id"] in failures:
                                failure = failures[row["id"]]
                                self.db.shipping(
                                    row,
                                    "failed",
                                    str(failure.get("message") or failure.get("code"))[:500],
                                )
                            else:
                                self.db.shipping(
                                    row, "unknown", "응답에 해당 주문 결과가 없습니다. 재조회 필요"
                                )
                    except Exception:
                        for row in batch:
                            self.db.shipping(
                                row,
                                "unknown",
                                "송신 응답을 확인하지 못했습니다. 자동 재송신하지 않습니다.",
                            )
                    await asyncio.sleep(0.55)
                try:
                    self.db.upsert(await api.details([r["id"] for r in sendable]))
                    self.reconcile()
                except Exception:
                    self.db.event(
                        "warning",
                        "송신 후 상세 재조회 실패. 주문 수집으로 반영 상태를 확인해 주세요.",
                    )
                results = [self.db.shipment(r["id"]) for r in valid]
                self.db.event(
                    "info",
                    f"송장 처리 {len(valid)}건 · "
                    f"반영확인 {sum(bool(r and r['state'] == 'success') for r in results)}건",
                )
                return results
            finally:
                self.busy = ""
                await api.close()
